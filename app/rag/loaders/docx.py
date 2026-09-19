from pathlib import Path

from docx import Document

from app.rag.loaders.base import (
    DocumentLoader,
    ExtractedDocument,
    ExtractedPage,
)
from app.utils.text_utils import clean_text


class DOCXLoader(DocumentLoader):
    def load(
        self,
        file_path: Path,
    ) -> ExtractedDocument:

        document = Document(file_path)

        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        full_text = clean_text("\n\n".join(paragraphs))

        pages = [
            ExtractedPage(
                page_number=None,
                text=full_text,
            )
        ]

        metadata = {
            "format": "docx",
            "paragraph_count": len(paragraphs),
        }

        return ExtractedDocument(
            pages=pages,
            metadata=metadata,
        )
